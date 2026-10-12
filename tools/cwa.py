"""Fetch the Continental Whippet Alliance's standings and title lists and
build the CWA feed (Whippet racing, beta).

The Continental Whippet Alliance (continentalwhippetalliance.com) races and
shows whippets and recognises both. It publishes:

- its grading guides (/grading-guides): the current one, a Google Sheet
  ("Updated: 10/4/2026") with every whippet on CWA's books by CWA number:
  grade, average, the last three meets (date, state, points), and progress
  toward each title (ARX points, NARX points, DPC points, High Combined
  wins, meets toward the TRP, Performance Racer points as of the last year
  end, completed meets); and an archived one with older whippets in the
  same layout (its meets carry a state but no date). Its "Ineligible" tab
  is not read;
- this year's standings on a page per year (/2026-toptens), as Datawrapper
  charts, each titled with its own date and meet count ("2026 Racing
  Standings as of 7/31/2026 (30 Race Meets)"): Racing (NARX points), High
  Combined (HC wins) and Conformation (not used here: this site covers
  coursing and racing). Each chart serves its data as a tab-separated CSV;
- the all-time racing lists, every whippet with 50 or more NARX points from
  1991 to the end of last year, also split into dogs and bitches, as three
  tabs of one published Google Sheet (/all-time-racing);
- a Google Sheet per title (/titles): ARX, NARX (levels), SRA (levels),
  DPC, DPCX, HC, HCX (levels), TRP and TRPX, each holder with the date or
  year earned.

Twenty-two downloads a week, a second or two apart: five pages from CWA's
own site (the grading guides page, the standings index, this year's
standings page, the titles page and the all-time page, read for their
links), then three charts from Datawrapper and fourteen sheets from Google.

Each whippet on a grading guide is keyed by its CWA number (a second
whippet the archive gives the same number gets "-2"); rows on the standings
and title lists join it by the registered name without its titles.
Whippets found only on those lists keep a slug of their name.

A snapshot per date (the newest of the guide's and the charts' dates) is
archived under data/cwa/snapshots/; later fetches with the same date
replace it. data/cwa.json (the current guide, standings, all-time list and
titles) and data/cwa-registry.json (the archived guide and list-only
whippets, loaded on demand) are built from the newest snapshot. The raw
downloads stay out of git (data/cwa/raw/): CWA's pages carry a contact
address. CWA was among the whippet racing bodies the site was asked to add.

Usage:
    python tools/cwa.py               # fetch, archive if changed, rebuild
    python tools/cwa.py --raw         # re-parse the files under data/cwa/raw
    python tools/cwa.py --offline     # rebuild from the snapshots alone
    python tools/cwa.py --force       # archive even if unchanged
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import time
from collections import Counter
from datetime import date, datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

from names import split_titles
from racing import (
    DATA, USER_AGENT, VOLATILE_KEYS, RacingParseError, fetch_bytes, load_snapshots,
    sha256, slug, today, write_json,
)

ORG = "cwa"
SITE_URL = "https://www.continentalwhippetalliance.com"
INDEX_URL = f"{SITE_URL}/toptens"
TITLES_URL = f"{SITE_URL}/titles"
GRADING_URL = f"{SITE_URL}/grading-guides"
SHEET_XLSX = "https://docs.google.com/spreadsheets/d/{id}/export?format=xlsx"
SHEET_VIEW = "https://docs.google.com/spreadsheets/d/{id}/edit"
ALL_TIME_URL = f"{SITE_URL}/all-time-racing"
SHEET_CSV = "https://docs.google.com/spreadsheets/d/{id}/export?format=csv"
RAW_DIR = DATA / "cwa" / "raw"
SNAPSHOT_DIR = DATA / "cwa" / "snapshots"
FEED = DATA / "cwa.json"
REGISTRY = DATA / "cwa-registry.json"

# Google answers a library User-Agent on a sheet export with a sign-in page.
BROWSER_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0 Safari/537.36"
)

# The titles, in CWA's own order, with their full names as CWA writes them.
TITLES = {
    "ARX": "Award of Racing Excellence",
    "NARX": "National Award of Racing Excellence",
    "SRA": "Superior Racing Award",
    "DPC": "Dual Purpose Champion",
    "DPCX": "Dual Purpose Champion Excellent",
    "HC": "High Combined",
    "HCX": "High Combined Excellent",
    "TRP": "Title of Racing Proficiency",
    "TRPX": "Title of Racing Proficiency Excellent",
}
RACING_TITLES = ("ARX", "NARX", "SRA", "TRP", "TRPX")
TITLE_COLUMNS = ["base", "title", "name", "earned", "year", "owners", "status", "dog"]
# The grading guides' columns, compared by their opening words.
GUIDE_FIELDS = ["CWA NO", "CALL NAME", "REGISTERED NAME", "OWNER", "GRADE", "AVG", "DPC", "HC",
                "ARX", "NARX", "TRP", "DQ", "DQ DATE", "LAST MEET", "", "2ND LAST MEET", "",
                "3RD LAST MEET", "", "PR", "YR LAST RUN"]
GUIDE_SHEETS = {"current": ("Current Grading Guide", "Grades"),
                "archive": ("Archived Grading Guide", "ARCHIVE")}
# A whippet as the feed packs it.
WHIPPET_COLUMNS = ["id", "cwa_no", "call_name", "registered_name", "name", "owner_raw", "grade",
                   "avg", "avg_computed", "dpc_points", "dpc_status", "hc", "arx_points",
                   "arx_status", "narx", "trp_meets", "trp_status", "dq", "dq_date", "meets",
                   "pr_points", "year_last_run", "never_raced", "completed_meets", "last_raced",
                   "on", "duplicate_of"]
LOC_RE = re.compile(r"^(\d{8})\s+([A-Z]{2})$")

STANDINGS_LINK_RE = re.compile(r"^(\d{4}) (Mid-Year|Year-End) Standings$")
CHART_TITLE_RE = re.compile(
    r"(\d{4}) (High Combined|Racing|Conformation) Standings as of "
    r"(\d{1,2})/(\d{1,2})/(\d{4})(?:\s*\((\d+) Race Meets\))?", re.IGNORECASE)
CHART_KINDS = {"high combined": "high_combined", "racing": "racing", "conformation": "conformation"}
CHART_HEADERS = {
    "racing": ["Place", "NARX Points", "Call Name", "Reg Name", "Owner"],
    "high_combined": ["Place", "HC Wins", "Call Name", "Reg Name", "Owner"],
    "conformation": ["Place", "Placement Points", "Call Name", "Reg Name", "Owner"],
}
ALL_TIME_LISTS = {"All Time Top Race Whippets": "all", "All Time Top Race Dogs": "dogs",
                  "All Time Top Race Bitches": "bitches"}
ALL_TIME_PERIOD_RE = re.compile(r"(\d{4}) through ([A-Z][a-z]+ \d{1,2}, \d{4})")
LEVEL_RE = re.compile(r"^(NARX\d?|SRA\d?)\b|^High Combined Excellent(?: (\d+))?$")
# Prefix titles CWA writes run together ("CanCH"); the shared vocabulary
# knows them only as separate words.
EXTRA_PREFIXES = {"canch", "amch", "amcanch", "intch"}


# ------------------------------------------------------------------- fetch

def _get(url: str, browser: bool = False) -> bytes:
    return fetch_bytes(url, user_agent=BROWSER_USER_AGENT if browser else USER_AGENT, timeout=60)


def _soup(raw: bytes) -> BeautifulSoup:
    return BeautifulSoup(raw.decode("utf-8", "replace"), "html.parser")


def standings_link(index_html: bytes) -> tuple[int, str, str]:
    """The newest standings page on the index: (year, kind, absolute URL)."""
    for anchor in _soup(index_html).find_all("a"):
        match = STANDINGS_LINK_RE.match(" ".join(anchor.get_text(" ").split()))
        if match and anchor.get("href"):
            return int(match.group(1)), match.group(2), urljoin(SITE_URL, anchor["href"])
    raise RacingParseError("The standings index lists no Mid-Year or Year-End standings")


def chart_frames(page_html: bytes) -> list[tuple[str, str]]:
    """(chart URL, chart title) for each Datawrapper chart on a standings page."""
    frames = []
    for frame in _soup(page_html).find_all("iframe"):
        src = frame.get("src") or ""
        if "datawrapper.dwcdn.net" in src:
            frames.append((src if src.endswith("/") else src + "/", frame.get("title") or ""))
    if not frames:
        raise RacingParseError("The standings page carries no Datawrapper charts")
    return frames


def chart_name(src: str) -> str:
    """'https://datawrapper.dwcdn.net/q52Ij/3/' -> 'q52Ij-3'."""
    match = re.search(r"dwcdn\.net/([A-Za-z0-9]+)/(\d+)/", src)
    if not match:
        raise RacingParseError(f"Unexpected chart address {src}")
    return f"{match.group(1)}-{match.group(2)}"


def title_sheets(titles_html: bytes) -> dict[str, str]:
    """{code: sheet id} from the titles page's links, matched by their text."""
    sheets = {}
    for anchor in _soup(titles_html).find_all("a"):
        href = anchor.get("href") or ""
        match = re.search(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)", href)
        text = " ".join(anchor.get_text(" ").split())
        code = re.search(r"\(([A-Z]+)\)\s*$", text)
        if match and code and code.group(1) in TITLES:
            sheets[code.group(1)] = match.group(1)
    missing = [code for code in TITLES if code not in sheets]
    if missing:
        raise RacingParseError(f"The titles page no longer links {missing}")
    return sheets


def guide_sheets(page_html: bytes) -> dict[str, str]:
    """{'current': sheet id, 'archive': sheet id} from the grading guides
    page's links, matched by their text."""
    found = {}
    for anchor in _soup(page_html).find_all("a"):
        text = " ".join(anchor.get_text(" ").split())
        match = re.search(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)", anchor.get("href") or "")
        for kind, (label, _sheet) in GUIDE_SHEETS.items():
            if match and text == label:
                found[kind] = match.group(1)
    missing = [kind for kind in GUIDE_SHEETS if kind not in found]
    if missing:
        raise RacingParseError(f"The grading guides page no longer links {missing}")
    return found


def all_time_tabs(page_html: bytes) -> list[tuple[str, str]]:
    """(gid, CSV URL) for each published tab on the all-time page."""
    tabs = []
    for frame in _soup(page_html).find_all("iframe"):
        src = frame.get("src") or ""
        match = re.search(r"(https://docs\.google\.com/spreadsheets/d/e/[A-Za-z0-9_-]+)/pubhtml\?gid=(\d+)", src)
        if match:
            tabs.append((match.group(2), f"{match.group(1)}/pub?gid={match.group(2)}&single=true&output=csv"))
    if len(tabs) < 3:
        raise RacingParseError(f"The all-time page shows {len(tabs)} sheet tabs, not three")
    return tabs


def fetch_all() -> dict[str, bytes]:
    """Every document, saved under data/cwa/raw/ by a fixed name."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    files: dict[str, bytes] = {}

    def keep(name: str, url: str, browser: bool = False) -> bytes:
        if files:
            time.sleep(1.5)
        data = _get(url, browser)
        (RAW_DIR / name).write_bytes(data)
        files[name] = data
        print(f"Fetched {url}  {len(data):,} bytes")
        return data

    for kind, sheet in guide_sheets(keep("grading.html", GRADING_URL)).items():
        keep(f"guide-{kind}.xlsx", SHEET_XLSX.format(id=sheet), browser=True)
    _, _, page_url = standings_link(keep("index.html", INDEX_URL))
    page = keep("standings.html", page_url)
    for src, _title in chart_frames(page):
        keep(f"chart-{chart_name(src)}.csv", src + "dataset.csv")
    titles = keep("titles.html", TITLES_URL)
    for code, sheet in title_sheets(titles).items():
        keep(f"title-{code.lower()}.csv", SHEET_CSV.format(id=sheet), browser=True)
    for gid, url in all_time_tabs(keep("all-time.html", ALL_TIME_URL)):
        keep(f"all-time-{gid}.csv", url, browser=True)
    return files


# ------------------------------------------------------------------- parse

def _date(month: str, day: str, year: str) -> str:
    return date(int(year), int(month), int(day)).isoformat()


def _number(text: str) -> float:
    try:
        return float(text.replace(",", ""))
    except ValueError as error:
        raise RacingParseError(f"Not a number: {text!r}") from error


def parse_chart(raw: bytes, title: str) -> tuple[str, dict]:
    """One standings chart: its kind and {label, as_of, meets, rows}."""
    match = CHART_TITLE_RE.search(title)
    if not match:
        raise RacingParseError(f"Unexpected chart title {title!r}")
    kind = CHART_KINDS[match.group(2).lower()]
    reader = list(csv.reader(io.StringIO(raw.decode("utf-8-sig", "replace")), delimiter="\t"))
    header = [cell.strip() for cell in reader[0]] if reader else []
    if header != CHART_HEADERS[kind]:
        raise RacingParseError(f"{match.group(2)} chart columns changed: {header}")
    rows = []
    for cells in reader[1:]:
        cells = [cell.strip() for cell in cells]
        if not any(cells):
            continue
        if len(cells) != 5 or not cells[0].isdigit():
            raise RacingParseError(f"{match.group(2)} chart row reads {cells}")
        rows.append({"place": int(cells[0]), "points": _number(cells[1]), "call_name": cells[2],
                     "registered_name": cells[3], "owners": cells[4]})
    return kind, {
        "label": " ".join(title.split()),
        "season": int(match.group(1)),
        "as_of": _date(match.group(3), match.group(4), match.group(5)),
        "meets": int(match.group(6)) if match.group(6) else None,
        "rows": rows,
    }


def parse_all_time(raw: bytes) -> tuple[str, dict]:
    """One all-time tab: which list it is and {criteria, period, through, rows}."""
    reader = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(raw.decode("utf-8-sig", "replace")))]
    heading = reader[0][0] if reader and reader[0] else ""
    if heading not in ALL_TIME_LISTS:
        raise RacingParseError(f"Unexpected all-time list {heading!r}")
    period = next((m for row in reader[:6] for m in [ALL_TIME_PERIOD_RE.search(row[0] if row else "")] if m), None)
    if not period:
        raise RacingParseError(f"{heading}: no period line")
    through = datetime.strptime(period.group(2), "%B %d, %Y").date().isoformat()
    try:
        start = next(i for i, row in enumerate(reader) if row[:4] == ["Place", "Name", "Owners", "NARX Points"])
    except StopIteration as error:
        raise RacingParseError(f"{heading}: the column headers changed") from error
    rows = []
    for row in reader[start + 1:]:
        if not any(row):
            continue
        if len(row) < 4 or not row[0].isdigit():
            raise RacingParseError(f"{heading}: row reads {row[:4]}")
        rows.append({"place": int(row[0]), "name": " ".join(row[1].split()), "owners": row[2],
                     "points": _number(row[3])})
    return ALL_TIME_LISTS[heading], {"criteria": reader[1][0] if len(reader) > 1 else "",
                                     "period": period.group(0), "through": through, "rows": rows}


def _level(cell: str) -> str | None:
    """A level heading inside a title sheet: 'NARX4 (60-74 NARX points)' ->
    'NARX4'; 'High Combined Excellent 2' -> 'HCX2'."""
    match = LEVEL_RE.match(cell)
    if not match:
        return None
    if match.group(1):
        return match.group(1)
    return f"HCX{match.group(2) or ''}"


def parse_title_sheet(code: str, raw: bytes) -> list[dict]:
    """A title sheet's holders. Three layouts: number, name, date (ARX, TRP,
    DPC, DPCX); number, name, year (HC, HCX); number, year, name, owners
    (NARX, SRA, TRPX). Levels head their own blocks in NARX, SRA and HCX."""
    reader = [[" ".join(cell.split()) for cell in row] for row in csv.reader(io.StringIO(raw.decode("utf-8-sig", "replace")))]
    if not reader or not reader[0] or not reader[0][0].endswith(f"({code})"):
        raise RacingParseError(f"The {code} sheet opens {reader[0][:1] if reader else None}")
    level = code
    rows = []
    for row in reader[2:]:
        cells = row + [""] * (4 - len(row))
        if not any(cells):
            continue
        heading = _level(cells[0])
        if heading:
            level = heading
            continue
        if not cells[0].isdigit():
            raise RacingParseError(f"{code}: unexpected row {cells[:4]}")
        number = int(cells[0])
        name, earned, year, owners, status = "", None, None, None, None
        # "Ineligible 5/13/2012": CWA lists the whippet in sequence with a
        # date, but marks it ineligible for the title. Kept, flagged, and
        # left out of the title counts.
        ineligible = re.fullmatch(r"Ineligible (\d{1,2})/(\d{1,2})/(\d{2}|\d{4})", cells[2], re.IGNORECASE)
        if ineligible:
            # A few are typed with a two-digit year ("9/24/17"); CWA's
            # records start in 1991.
            month, day, yr = ineligible.groups()
            if len(yr) == 2:
                yr = str((1900 if int(yr) >= 90 else 2000) + int(yr))
            cells[2], status = f"{month}/{day}/{yr}", "ineligible"
        if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", cells[2]):
            name = cells[1]
            month, day, yr = cells[2].split("/")
            earned, year = _date(month, day, yr), int(yr)
        elif re.fullmatch(r"(19|20)\d\d", cells[2]):
            name, year = cells[1], int(cells[2])
        elif cells[2]:
            name, owners = cells[2], cells[3] or None
            # The year cell may carry the flag instead: "2013 Ineligible".
            flagged = re.fullmatch(r"((?:19|20)\d\d)\s+Ineligible", cells[1], re.IGNORECASE)
            if flagged:
                year, status = int(flagged.group(1)), "ineligible"
            elif re.fullmatch(r"(19|20)\d\d", cells[1]):
                year = int(cells[1])
            elif cells[1]:
                raise RacingParseError(f"{code}: year reads {cells[1]!r}")
        elif cells[1]:
            name = cells[1]
        if not name:
            continue   # a numbered slot with nobody in it
        rows.append({"number": number, "title": level, "name": name, "earned": earned,
                     "year": year, "owners": owners, "status": status})
    if not rows:
        raise RacingParseError(f"The {code} sheet lists nobody")
    return rows


def _head(value) -> str:
    return " ".join(str(value or "").replace("_", " ").split()).upper()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _num(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def _points_or_mark(value) -> tuple[float | None, str | None]:
    """A progress cell holds points (or meets) toward a title, or the mark
    CWA writes once it is earned ("ARX", "TRP30", "DPC", "CH")."""
    number = _num(value)
    if number is not None:
        return number, None
    return None, _text(value) or None


def parse_guide(raw: bytes, kind: str) -> dict:
    """One grading guide: its date (the current one only) and its whippets."""
    import openpyxl

    label, sheet = GUIDE_SHEETS[kind]
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if sheet not in book.sheetnames:
        raise RacingParseError(f"The {label} has no {sheet!r} sheet: {book.sheetnames}")
    rows = [list(row) for row in book[sheet].iter_rows(values_only=True)]
    try:
        top = next(i for i, row in enumerate(rows[:6]) if row and _head(row[0]) == "CWA NO")
    except StopIteration as error:
        raise RacingParseError(f"The {label} has no CWA NO column") from error
    header = [_head(cell) for cell in rows[top]] + [""] * 30
    for position, want in enumerate(GUIDE_FIELDS):
        if want and not header[position].startswith(want):
            raise RacingParseError(f"{label}: column {position + 1} reads {header[position]!r}, not {want!r}")
    updated = None
    if kind == "current":
        stamp = re.search(r"Updated:\s*(\d{1,2})/(\d{1,2})/(\d{4})", _text(rows[0][0] if rows[0] else ""))
        if not stamp:
            raise RacingParseError("The current grading guide has no 'Updated' line")
        updated = _date(*stamp.groups())
    as_of = re.search(r"AS OF (\d{1,2})/(\d{1,2})/(\d{4})", header[19])
    completed_col = next((i for i, head in enumerate(header) if head.startswith("COMPLETED MEETS")), None)
    whippets = []
    for row in rows[top + 2:]:
        row = list(row) + [None] * 30
        call, registered = _text(row[1]), _text(row[2])
        if not call and not registered:
            continue   # CWA's blank template rows
        meets = []
        for loc_col, pts_col in ((13, 14), (15, 16), (17, 18)):
            place, points = _text(row[loc_col]), _num(row[pts_col])
            if not place and points is None:
                continue
            dated = LOC_RE.match(place)
            if dated:
                day = dated.group(1)
                meets.append([_date(day[4:6], day[6:8], day[0:4]), dated.group(2), points])
            else:
                # A state alone (the archive dates no meet), or whatever CWA
                # typed there ("14.5" once): kept as written.
                meets.append([None, place or None, points])
        dq_text = _text(row[12])
        dq_dated = LOC_RE.match(dq_text)
        last_run = _text(row[20])
        dpc_points, dpc_status = _points_or_mark(row[6])
        arx_points, arx_status = _points_or_mark(row[8])
        trp_meets, trp_status = _points_or_mark(row[10])
        completed = _num(row[completed_col]) if completed_col is not None else None
        whippets.append({
            "cwa_no": _text(row[0]) or None,
            "call_name": call,
            "registered_name": registered,
            "owner_raw": _text(row[3]),
            "grade": _text(row[4]) or None,
            "avg": _num(row[5]),
            "dpc_points": dpc_points, "dpc_status": dpc_status,
            "hc": _num(row[7]),
            "arx_points": arx_points, "arx_status": arx_status,
            "narx": _num(row[9]),
            "trp_meets": trp_meets, "trp_status": trp_status,
            "dq": _num(row[11]),
            "dq_date": ([_date(dq_dated.group(1)[4:6], dq_dated.group(1)[6:8], dq_dated.group(1)[0:4]),
                         dq_dated.group(2)] if dq_dated else ([None, dq_text] if dq_text else None)),
            "meets": meets,
            "pr_points": _num(row[19]),
            "year_last_run": int(last_run) if last_run.isdigit() else None,
            "never_raced": last_run.upper() == "NR",
            "completed_meets": int(completed) if completed is not None else None,
        })
    if not whippets:
        raise RacingParseError(f"The {label} lists no whippets")
    return {"updated": updated, "pr_as_of": _date(*as_of.groups()) if as_of else None,
            "rows": whippets}


def parse_all(files: dict[str, bytes]) -> dict:
    season, kind, page_url = standings_link(files["index.html"])
    standings = {}
    for src, title in chart_frames(files["standings.html"]):
        chart_kind, chart = parse_chart(files[f"chart-{chart_name(src)}.csv"], title)
        if chart_kind == "conformation":
            continue   # conformation alone is outside this site's scope
        chart["chart_url"] = src
        standings[chart_kind] = chart
    for needed in ("racing", "high_combined"):
        if needed not in standings:
            raise RacingParseError(f"The standings page has no {needed.replace('_', ' ')} chart")
    titles = {}
    for code, sheet in title_sheets(files["titles.html"]).items():
        titles[code] = {"sheet": sheet, "rows": parse_title_sheet(code, files[f"title-{code.lower()}.csv"])}
    all_time = {}
    for gid, _url in all_time_tabs(files["all-time.html"]):
        which, listing = parse_all_time(files[f"all-time-{gid}.csv"])
        listing["gid"] = gid
        all_time[which] = listing
    if set(all_time) != set(ALL_TIME_LISTS.values()):
        raise RacingParseError(f"All-time lists found: {sorted(all_time)}")
    sheets = guide_sheets(files["grading.html"])
    guide = {kind: {**parse_guide(files[f"guide-{kind}.xlsx"], kind), "sheet": sheet}
             for kind, sheet in sheets.items()}
    return {
        "org": ORG,
        "guide_date": max([chart["as_of"] for chart in standings.values()] + [guide["current"]["updated"]]),
        "guide_date_source": "guide",
        "season": season,
        "standings_kind": kind,
        "source_url": page_url,
        "sha256": sha256(files["standings.html"]),
        "fetched": today(),
        "standings": standings,
        "all_time": all_time,
        "titles": titles,
        "guide": guide,
    }


def write_snapshot(snapshot: dict, force: bool = False) -> None:
    """Archive the parse unless it matches the newest snapshot (fetch-time
    keys aside); a change under the same standings date replaces that
    date's file, as titles are added between monthly charts."""
    SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)

    def stable(snap: dict) -> str:
        return json.dumps({k: v for k, v in snap.items() if k not in VOLATILE_KEYS},
                          sort_keys=True, ensure_ascii=False)

    existing = sorted(SNAPSHOT_DIR.glob("*.json"))
    if existing and not force:
        newest = json.loads(existing[-1].read_text(encoding="utf-8"))
        if stable(newest) == stable(snapshot):
            print(f"Unchanged since {existing[-1].name} - nothing archived.")
            return
    target = SNAPSHOT_DIR / f"{snapshot['guide_date']}.json"
    target.write_text(json.dumps(snapshot, ensure_ascii=False, indent=1) + "\n",
                      encoding="utf-8", newline="\n")
    print(f"Archived {target.name}")


# ------------------------------------------------------------------- build

def core_name(name: str) -> str:
    """The registered name without its titles, for grouping one whippet's
    rows across lists that print different titles after its name."""
    words = name.split()
    while words and words[0].lower().strip(".") in EXTRA_PREFIXES:
        words = words[1:]
    _prefix, core, _suffix = split_titles(" ".join(words))
    return core


def dog_id(name: str) -> str:
    return slug(core_name(name)) or slug(name)


def build(snapshots: list[dict]) -> tuple[dict, dict]:
    current = snapshots[-1]
    season = current["season"]
    guide = current["guide"]

    # ---- whippets: one per row of the grading guides, by CWA number
    whippets: dict[str, dict] = {}
    by_core: dict[str, str] = {}
    for where in ("current", "archive"):
        for row in guide[where]["rows"]:
            number = row["cwa_no"] if row["cwa_no"] and re.fullmatch(r"[0-9A-Za-z]+", row["cwa_no"]) else None
            base = number or dog_id(row["registered_name"] or row["call_name"])
            wid, n = base, 1
            while wid in whippets:
                n += 1
                wid = f"{base}-{n}"
            listed = [meet[2] for meet in row["meets"] if meet[2] is not None]
            computed = round(sum(listed) / len(listed), 6) if listed else 0.0
            whippets[wid] = {
                **row, "id": wid, "on": where, "duplicate_of": base if wid != base else None,
                "name": core_name(row["registered_name"]) or row["registered_name"] or row["call_name"],
                "avg_computed": computed,
                "last_raced": max((meet[0] for meet in row["meets"] if meet[0]), default=None),
            }
            key = dog_id(row["registered_name"]) if row["registered_name"] else None
            if key and key not in by_core:
                by_core[key] = wid   # the current guide first, so it wins

    # ---- the lists, joined by the registered name without titles; a whippet
    #      on no guide keeps a slug of its name
    list_only: dict[str, dict] = {}
    current_calls = Counter(row["call_name"].lower() for row in whippets.values() if row["on"] == "current")

    def seen(name: str, call_name: str | None = None) -> str:
        key = dog_id(name)
        if key in by_core:
            return by_core[key]
        if call_name and current_calls[call_name.lower()] == 1:
            return next(w["id"] for w in whippets.values()
                        if w["on"] == "current" and w["call_name"].lower() == call_name.lower())
        key = key if key not in whippets else f"l-{key}"
        entry = list_only.setdefault(key, {"id": key, "name": core_name(name) or name,
                                           "call_name": call_name or "", "on": "lists"})
        if call_name and not entry["call_name"]:
            entry["call_name"] = call_name
        return key

    standings = {}
    for kind, chart in current["standings"].items():
        rows = [{**row, "dog": seen(row["registered_name"], row["call_name"])} for row in chart["rows"]]
        standings[kind] = {**{k: v for k, v in chart.items() if k != "rows"}, "rows": rows}

    lists = current["all_time"]
    sex = {}
    for which, label in (("dogs", "dog"), ("bitches", "bitch")):
        for row in lists[which]["rows"]:
            sex[dog_id(row["name"])] = label
    all_time = {
        "criteria": lists["all"]["criteria"], "period": lists["all"]["period"],
        "through": lists["all"]["through"],
        "rows": [{**row, "sex": sex.get(dog_id(row["name"])), "dog": seen(row["name"])}
                 for row in lists["all"]["rows"]],
    }

    titles = []
    for code, sheet in current["titles"].items():
        for row in sheet["rows"]:
            titles.append({"base": code, **{k: row[k] for k in ("title", "name", "earned", "year", "owners", "status")},
                           "dog": seen(row["name"])})
    # Newest first: by year, then by date where the sheet gives one.
    titles.sort(key=lambda row: (row["base"], row["name"]))
    titles.sort(key=lambda row: (row["year"] or 0, row["earned"] or ""), reverse=True)

    on_guide = [w for w in whippets.values() if w["on"] == "current"]
    archived = [w for w in whippets.values() if w["on"] == "archive"]
    awarded = [row for row in titles if row["status"] is None]
    this_year = [row for row in awarded if row["year"] == season]
    racing = standings["racing"]
    leader = min(racing["rows"], key=lambda row: (row["place"], row["call_name"]), default=None)
    with_meets = [w for w in on_guide if w["meets"] and w["avg"] is not None]
    stats = {
        "season": season,
        "guide_updated": guide["current"]["updated"],
        "pr_as_of": guide["current"]["pr_as_of"],
        "whippets_current": len(on_guide),
        "raced_this_year": sum(w["year_last_run"] == season for w in on_guide),
        "never_raced": sum(w["never_raced"] for w in on_guide),
        "grades_current": dict(sorted(Counter(w["grade"] or "" for w in on_guide).items())),
        "avg_agreement": round(sum(abs(w["avg"] - w["avg_computed"]) <= 0.01 for w in with_meets)
                               / max(1, len(with_meets)), 4),
        "archive_listed": len(archived),
        "archive_shared_numbers": sum(w["duplicate_of"] is not None for w in archived),
        "list_only": len(list_only),
        "racing_as_of": racing["as_of"],
        "racing_meets": racing["meets"],
        "racing_listed": len(racing["rows"]),
        "high_combined_as_of": standings["high_combined"]["as_of"],
        "high_combined_meets": standings["high_combined"]["meets"],
        "high_combined_listed": len(standings["high_combined"]["rows"]),
        "titles_total": len(awarded),
        "titles_ineligible": len(titles) - len(awarded),
        "titles_this_year": len(this_year),
        "titles_this_year_by_base": dict(Counter(row["base"] for row in this_year)),
        "title_holders": len({row["dog"] for row in awarded}),
        "arx_total": sum(row["base"] == "ARX" for row in awarded),
        "trp_total": sum(row["base"] == "TRP" for row in awarded),
        "dpc_total": sum(row["base"] == "DPC" for row in awarded),
        "all_time_listed": len(all_time["rows"]),
        "latest_title": max((row["earned"] for row in awarded if row["earned"]), default=None),
        "leader": ({"dog": leader["dog"], "call_name": leader["call_name"], "points": leader["points"]}
                   if leader else None),
    }

    def pack(row: dict) -> list:
        return [row.get(column) for column in WHIPPET_COLUMNS]

    feed = {
        "org": ORG,
        "name": "CWA",
        "program": "whippet racing",
        "season": season,
        "guide_date": current["guide_date"],
        "guide_date_source": current["guide_date_source"],
        "generated": today(),
        "fetched": current["fetched"],
        "site_url": SITE_URL,
        "guide_page": GRADING_URL,
        "guide_url": SHEET_VIEW.format(id=guide["current"]["sheet"]),
        "archive_url": SHEET_VIEW.format(id=guide["archive"]["sheet"]),
        "standings_url": current["source_url"],
        "titles_url": TITLES_URL,
        "all_time_url": ALL_TIME_URL,
        "snapshots": [snap["guide_date"] for snap in snapshots],
        "stats": stats,
        "title_names": TITLES,
        "racing_titles": list(RACING_TITLES),
        "standings": standings,
        "all_time": all_time,
        # The current guide's whippets, packed by whippet_columns; the archive
        # and list-only whippets are in data/cwa-registry.json.
        "whippet_columns": WHIPPET_COLUMNS,
        "whippets": [pack(w) for w in sorted(on_guide, key=lambda w: (w["call_name"].lower(), w["id"]))],
        # Title rows packed as arrays (title_columns); dog is a whippet id.
        "title_columns": TITLE_COLUMNS,
        "titles": [[row["base"], row["title"], row["name"], row["earned"], row["year"],
                    row["owners"], row["status"], row["dog"]] for row in titles],
        "notes": [
            "Whippets are CWA's grading guides as written, keyed by CWA number; avg_computed "
            "is the mean of the listed meets (CWA rule: the last three completed meets).",
            "Standings are CWA's own charts, each with its own date and meet count; "
            "conformation standings are left out. Places are CWA's: ties share a place.",
            "Titles are CWA's title sheets as written, joined to a whippet by its "
            "registered name without titles; names are shown as each list prints them.",
        ],
    }
    registry_rows = sorted(archived, key=lambda w: ((w["call_name"] or w["name"]).lower(), w["id"]))
    registry_rows += sorted(list_only.values(), key=lambda w: (w["name"].lower(), w["id"]))
    registry = {"org": ORG, "guide_date": current["guide_date"], "season": season,
                "columns": WHIPPET_COLUMNS, "rows": [pack(row) for row in registry_rows]}
    return feed, registry


# -------------------------------------------------------------------- main

def read_raw() -> dict[str, bytes]:
    return {path.name: path.read_bytes() for path in RAW_DIR.iterdir() if path.is_file()}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true", help="archive even if unchanged")
    parser.add_argument("--offline", action="store_true",
                        help="rebuild the feed from archived snapshots without fetching")
    parser.add_argument("--raw", action="store_true",
                        help="re-parse the files already under data/cwa/raw instead of fetching")
    args = parser.parse_args()

    if not args.offline:
        try:
            files = read_raw() if args.raw else fetch_all()
            snapshot = parse_all(files)
            if args.raw:
                snapshot["fetched"] = date.fromtimestamp(
                    (RAW_DIR / "standings.html").stat().st_mtime).isoformat()
            write_snapshot(snapshot, force=args.force)
        except requests.RequestException as error:
            if not list(SNAPSHOT_DIR.glob("*.json")):
                print(f"Could not fetch from CWA and no snapshot exists: {error}", file=sys.stderr)
                return 1
            print(f"Could not fetch from CWA; rebuilding from archived snapshots. {error}",
                  file=sys.stderr)

    snapshots = load_snapshots(SNAPSHOT_DIR)
    feed, registry = build(snapshots)
    size = write_json(FEED, feed)
    registry_size = write_json(REGISTRY, registry)
    stats = feed["stats"]
    print(
        f"CWA standings {feed['guide_date']}: racing as of {stats['racing_as_of']} "
        f"({stats['racing_listed']} listed, {stats['racing_meets']} meets), High Combined as of "
        f"{stats['high_combined_as_of']} ({stats['high_combined_listed']} listed); "
        f"{stats['titles_total']:,} titles ({stats['titles_this_year']} in {stats['season']}), "
        f"{stats['all_time_listed']} on the all-time list"
    )
    print(f"Grading guide {stats['guide_updated']}: {stats['whippets_current']} whippets, "
          f"{stats['raced_this_year']} racing this year, average agreement {stats['avg_agreement']:.1%}; "
          f"archive {stats['archive_listed']:,} ({stats['archive_shared_numbers']} on a shared number), "
          f"{stats['list_only']:,} on the lists only")
    print(f"Wrote {FEED.name} ({size:,} bytes) and {REGISTRY.name} ({registry_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
