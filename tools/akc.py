"""Fetch the AKC lure coursing entries workbook and build data/akc.json.

The workbook (a shared Google Sheet, AKC_Lure_Coursing_Entries_2026) is
compiled by hand from AKC's Event Search: one row per AKC lure coursing
event (type "Lure Coursing Tests & Trials") with the entry and competitor
counts from that event's results page, plus Summary, ranking and Notes tabs
whose figures are formulas over the rows. It is re-collected now and then;
the Summary tab says when ("Collected October 9, 2026") and the period the
rows cover.

Every figure the site shows is recomputed here from the rows and must equal
the sheet's own formulas, tab by tab; a difference stops the build rather
than publishing a number the sheet does not carry.

Usage:
    python tools/akc.py                 # fetch, archive if changed, rebuild
    python tools/akc.py --file PATH     # parse a local export instead
    python tools/akc.py --offline       # rebuild from the newest archived export
    python tools/akc.py --date YYYY-MM-DD   # override the collected date
    python tools/akc.py --force         # archive even if unchanged
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
import requests

from events import US_STATES
from racing import DATA, RacingParseError, cell_num, cell_str, fetch_bytes, sha256, today, write_json

ORG = "akc"
SHEET_ID = "1N9XyQU5sLPxFxCoHYr52GkMyCeQYKt_-QG8JVuH88R4"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?usp=sharing"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
SOURCE_NAME = "AKC Event Search"
SOURCE_URL = "https://webapps.akc.org/event-search/#/results"
EVENT_URL = (
    "https://www.apps.akc.org/apps/events/search/index_results.cfm"
    "?action=plan&event_number={number}"
)
RAW_DIR = DATA / "akc" / "raw"
FEED = DATA / "akc.json"

# Google's export endpoint answers a plain library User-Agent with a sign-in
# page; a browser-style one gets the workbook (the same exception aok9.py
# makes to the identify-the-project convention in racing.USER_AGENT).
EXPORT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)

TAB_TRIALS = "Trials"
TAB_SUMMARY = "Summary"
TAB_TRIAL_RANKING = "Trial ranking"
TAB_CLUB_RANKING = "Club ranking"
TAB_STATE_RANKING = "State ranking"
TAB_NOTES = "Notes"

HEADERS = [
    "Date", "Day", "Club", "City", "State", "AKC event #", "Scope",
    "LC Trial entries", "LC Trial competitors", "LC Test entries (JC/QC)",
    "LC Test competitors", "Total entries", "Results posted", "AKC event page",
]
SUMMARY_LABELS = {
    "Events listed": "events",
    "Events with results posted": "with_results",
    "Events without results posted yet": "without_results",
    "LC Trial entries": "trial_entries",
    "LC Trial competitors": "trial_competitors",
    "LC Test entries (JC/QC)": "test_entries",
    "LC Test competitors": "test_competitors",
    "Total entries (trial + test)": "total_entries",
    "Events that held an LC Trial": "events_with_trial",
    "Average LC Trial entries per trial": "avg_trial_entries",
    "Largest LC Trial entry": "largest_trial",
    "Smallest LC Trial entry": "smallest_trial",
}
MONTH_LABELS = ["January", "February", "March", "April", "May", "June", "July",
                "August", "September", "October", "November", "December"]
EVENT_NUMBER_RE = re.compile(r"^20\d{8}$")
PERIOD_RE = re.compile(r"(\w+ \d{1,2}) to (\w+ \d{1,2}), (\d{4})")
COLLECTED_RE = re.compile(r"Collected (\w+ \d{1,2}, \d{4})")

# The workbook is committed, so it is checked for the contact details the
# other raw workbooks are kept out of git for.
PRIVATE_PATTERNS = (
    ("email address", re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")),
    ("phone number", re.compile(r"\(?\d{3}\)?[-. ]\d{3}[-. ]\d{4}")),
    ("street address", re.compile(
        r"\b\d{1,5}\s+\w+\s+(St|Street|Ave|Avenue|Rd|Road|Dr|Drive|Ln|Lane|Blvd|Way|Ct|Court)\b",
        re.IGNORECASE)),
)


def _norm(value) -> str:
    return re.sub(r"\s+", "", cell_str(value)).lower()


def _int(value) -> int | None:
    number = cell_num(value)
    if number is None:
        return None
    if number != int(number):
        raise RacingParseError(f"expected a whole number, got {value!r}")
    return int(number)


def _iso(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = cell_str(value)
    for fmt in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise RacingParseError(f"unreadable date {value!r}")


# ------------------------------------------------------------------- fetch
def fetch() -> bytes:
    raw = fetch_bytes(EXPORT_URL, user_agent=EXPORT_USER_AGENT)
    if not raw.startswith(b"PK"):
        raise RacingParseError(
            "The export did not return a workbook (no zip signature); Google "
            "may have answered with a sign-in page. Is the sheet shared by link?"
        )
    print(f"Fetched the AKC entries workbook  {len(raw):,} bytes")
    return raw


# ------------------------------------------------------------------- parse
def read_trials(ws) -> list[dict]:
    rows = list(ws.iter_rows(min_row=1))
    if not rows:
        raise RacingParseError("the Trials tab is empty")
    header = [_norm(c.value) for c in rows[0][:len(HEADERS)]]
    if header != [_norm(h) for h in HEADERS]:
        raise RacingParseError(f"Trials headers changed: {[c.value for c in rows[0][:len(HEADERS)]]}")
    events: list[dict] = []
    seen: set[str] = set()
    for cells in rows[1:]:
        values = [c.value for c in cells[:len(HEADERS)]]
        if values[0] is None or cell_str(values[0]) == "":
            break
        when = _iso(values[0])
        day = cell_str(values[1])
        weekday = date.fromisoformat(when).strftime("%a")
        if day != weekday:
            raise RacingParseError(f"{when}: the Day column says {day!r}, the date is a {weekday}")
        number = cell_str(values[5])
        if not EVENT_NUMBER_RE.match(number):
            raise RacingParseError(f"{when}: unexpected AKC event number {values[5]!r}")
        if number in seen:
            raise RacingParseError(f"event number {number} appears twice")
        seen.add(number)
        state = cell_str(values[4])
        if state not in US_STATES:
            raise RacingParseError(f"{when} {number}: unknown state {state!r}")
        scope = cell_str(values[6]) or None
        posted_text = cell_str(values[12]).lower()
        if posted_text not in ("yes", "no"):
            raise RacingParseError(f"{number}: Results posted is {values[12]!r}, not Yes/No")
        link = cells[13].hyperlink.target if cells[13].hyperlink else None
        url = link or EVENT_URL.format(number=number)
        if number not in url:
            raise RacingParseError(f"{number}: the event link {url} is for another event")
        event = {
            "date": when,
            "day": day,
            "club": cell_str(values[2]),
            "city": cell_str(values[3]),
            "state": state,
            "event_number": number,
            "scope": scope,
            "specialty": scope is not None and scope.lower() != "all-breed",
            "trial_entries": _int(values[7]),
            "trial_competitors": _int(values[8]),
            "test_entries": _int(values[9]),
            "test_competitors": _int(values[10]),
            "total_entries": _int(values[11]),
            "results_posted": posted_text == "yes",
            "url": url,
        }
        if not event["club"] or not event["city"]:
            raise RacingParseError(f"{number}: club or city missing")
        if not event["results_posted"]:
            if any(event[k] is not None for k in ("trial_entries", "trial_competitors",
                                                   "test_entries", "test_competitors", "total_entries")):
                raise RacingParseError(f"{number}: counts on an event whose results are not posted")
        else:
            expected_total = (event["trial_entries"] or 0) + (event["test_entries"] or 0)
            if (event["total_entries"] or 0) != expected_total:
                raise RacingParseError(
                    f"{number}: total {event['total_entries']} is not trial + test = {expected_total}")
        events.append(event)
    if not events:
        raise RacingParseError("no event rows found")
    events.sort(key=lambda e: (e["date"], e["event_number"]))
    return events


def read_summary(ws) -> dict:
    title = cell_str(ws.cell(1, 1).value)
    source_line = cell_str(ws.cell(2, 1).value)
    period = PERIOD_RE.search(title)
    collected = COLLECTED_RE.search(source_line)
    if not period:
        raise RacingParseError(f"Summary A1 carries no period: {title!r}")
    if not collected:
        raise RacingParseError(f"Summary A2 carries no 'Collected' date: {source_line!r}")
    year = period.group(3)
    summary = {
        "period_start": _iso(f"{period.group(1)}, {year}"),
        "period_end": _iso(f"{period.group(2)}, {year}"),
        "collected": _iso(collected.group(1)),
        "figures": {},
        "by_month": {},
    }
    labels = {_norm(label): key for label, key in SUMMARY_LABELS.items()}
    month_header_row = None
    for row in range(3, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label is None:
            continue
        key = labels.get(_norm(label))
        if key:
            summary["figures"][key] = cell_num(ws.cell(row, 2).value)
        elif _norm(label) == "month":
            month_header_row = row
            break
    missing = [k for k in SUMMARY_LABELS.values() if k not in summary["figures"]]
    if missing or month_header_row is None:
        raise RacingParseError(f"Summary tab changed shape: missing {missing or 'the by-month table'}")
    header = [_norm(ws.cell(month_header_row, c).value) for c in range(1, 8)]
    expected = [_norm(h) for h in ("Month", "Events", "With results", "LC Trial entries",
                                   "LC Test entries", "Total entries", "Avg trial entries")]
    if header != expected:
        raise RacingParseError(f"By-month header changed: {header}")
    for row in range(month_header_row + 1, ws.max_row + 1):
        month = ws.cell(row, 1).value
        if month is None:
            break
        if isinstance(month, str) and _norm(month) == "total":
            summary["by_month_total"] = {
                "events": _int(ws.cell(row, 2).value) or 0,
                "trial_entries": _int(ws.cell(row, 4).value) or 0,
                "test_entries": _int(ws.cell(row, 5).value) or 0,
            }
            break
        key = _iso(month)[:7]
        summary["by_month"][key] = {
            "events": _int(ws.cell(row, 2).value) or 0,
            "with_results": _int(ws.cell(row, 3).value) or 0,
            "trial_entries": _int(ws.cell(row, 4).value) or 0,
            "test_entries": _int(ws.cell(row, 5).value) or 0,
            "total_entries": _int(ws.cell(row, 6).value) or 0,
            "avg_trial_entries": cell_num(ws.cell(row, 7).value),
        }
    return summary


def _ranking_rows(ws, expected_header: list[str]) -> list[list]:
    header_row = None
    for row in range(1, min(ws.max_row, 10) + 1):
        if _norm(ws.cell(row, 1).value) == "rank":
            header_row = row
            break
    if header_row is None:
        raise RacingParseError(f"{ws.title}: no Rank header")
    header = [_norm(ws.cell(header_row, c).value) for c in range(1, len(expected_header) + 1)]
    if header != [_norm(h) for h in expected_header]:
        raise RacingParseError(f"{ws.title}: header changed: {header}")
    rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row, c).value for c in range(1, len(expected_header) + 1)]
        if all(v is None or cell_str(v) == "" for v in values):
            continue
        rows.append(values)
    return rows


def read_rankings(wb) -> dict:
    trials = []
    for rank, when, club, city, state, entries, competitors, tests in _ranking_rows(
            wb[TAB_TRIAL_RANKING], ["Rank", "Date", "Club", "City", "State", "LC Trial entries",
                                    "LC Trial competitors", "LC Test entries"]):
        trials.append({
            "rank": _int(rank), "date": _iso(when), "club": cell_str(club), "city": cell_str(city),
            "state": cell_str(state), "trial_entries": _int(entries),
            "trial_competitors": _int(competitors), "test_entries": _int(tests),
        })
    clubs = []
    for rank, club, held, entries, avg, largest, tests in _ranking_rows(
            wb[TAB_CLUB_RANKING], ["Rank", "Club", "Trials held", "LC Trial entries", "Avg per trial",
                                   "Largest trial", "LC Test entries"]):
        clubs.append({
            "rank": _int(rank), "club": cell_str(club), "trials": _int(held),
            "trial_entries": _int(entries), "avg_trial_entries": cell_num(avg),
            "largest_trial": _int(largest), "test_entries": _int(tests),
        })
    states = []
    for rank, state, held, entries, avg, largest, tests in _ranking_rows(
            wb[TAB_STATE_RANKING], ["Rank", "State", "Trials held", "LC Trial entries", "Avg per trial",
                                    "Largest trial", "LC Test entries"]):
        states.append({
            "rank": _int(rank), "state": cell_str(state), "trials": _int(held),
            "trial_entries": _int(entries), "avg_trial_entries": cell_num(avg),
            "largest_trial": _int(largest), "test_entries": _int(tests),
        })
    return {"trials": trials, "clubs": clubs, "states": states}


# Notes about the workbook's own mechanics (its formulas, how its tabs
# re-sort) are for whoever edits the sheet, not for the page.
WORKBOOK_NOTE_RE = re.compile(r"formula|re-sort", re.IGNORECASE)


def read_notes(ws) -> list[str]:
    notes = [cell_str(ws.cell(row, 1).value) for row in range(1, ws.max_row + 1)]
    notes = [n for n in notes if n]
    if notes and notes[0].lower().startswith("how to read"):
        notes = notes[1:]
    return [n for n in notes if not WORKBOOK_NOTE_RE.search(n)]


def assert_no_personal_data(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if not isinstance(value, str):
                    continue
                for label, pattern in PRIVATE_PATTERNS:
                    if pattern.search(value):
                        raise RacingParseError(
                            f"{ws.title}: a cell looks like a {label}; the workbook is "
                            f"committed to the repo, so this cannot be published")


def parse_workbook(raw: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    for tab in (TAB_TRIALS, TAB_SUMMARY, TAB_TRIAL_RANKING, TAB_CLUB_RANKING, TAB_STATE_RANKING, TAB_NOTES):
        if tab not in wb.sheetnames:
            raise RacingParseError(f"the workbook has no {tab!r} tab: {wb.sheetnames}")
    assert_no_personal_data(wb)
    return {
        "events": read_trials(wb[TAB_TRIALS]),
        "summary": read_summary(wb[TAB_SUMMARY]),
        "rankings": read_rankings(wb),
        "notes": read_notes(wb[TAB_NOTES]),
        "modified": wb.properties.modified.date().isoformat() if wb.properties.modified else None,
    }


def content_key(parsed: dict) -> str:
    return json.dumps({k: parsed[k] for k in ("events", "summary", "rankings", "notes")},
                      sort_keys=True, ensure_ascii=False)


# ------------------------------------------------------------------- build
def competition_rank(rows: list[dict], field: str) -> None:
    """RANK() as the sheet uses it: ties share a rank, the next rank skips.
    Rows without a value stay unranked."""
    ordered = sorted((r for r in rows if r.get(field) is not None), key=lambda r: -r[field])
    previous = None
    rank = 0
    for position, row in enumerate(ordered, 1):
        if row[field] != previous:
            rank, previous = position, row[field]
        row["rank"] = rank
    for row in rows:
        row.setdefault("rank", None)


def _round(value: float | None, places: int = 3) -> float | None:
    return None if value is None else round(value, places)


def build(parsed: dict, source_file: str, digest: str, collected: str | None = None) -> dict:
    events = parsed["events"]
    summary = parsed["summary"]
    collected = collected or summary["collected"]
    period = {"start": summary["period_start"], "end": summary["period_end"]}
    season = int(period["start"][:4])

    posted = [e for e in events if e["results_posted"]]
    held = [e for e in posted if e["trial_entries"] is not None]
    trial_sum = sum(e["trial_entries"] for e in held)
    stats = {
        "events": len(events),
        "with_results": len(posted),
        "without_results": len(events) - len(posted),
        "events_with_trial": len(held),
        "trial_entries": trial_sum,
        "trial_competitors": sum(e["trial_competitors"] or 0 for e in posted),
        "test_entries": sum(e["test_entries"] or 0 for e in posted),
        "test_competitors": sum(e["test_competitors"] or 0 for e in posted),
        "total_entries": sum(e["total_entries"] or 0 for e in posted),
        "avg_trial_entries": _round(trial_sum / len(held)) if held else None,
        "largest_trial": max((e["trial_entries"] for e in held), default=None),
        "smallest_trial": min((e["trial_entries"] for e in held), default=None),
        "clubs": len({e["club"] for e in events}),
        "states": len({e["state"] for e in events}),
        "specialties": sum(e["specialty"] for e in events),
    }

    by_month = []
    start, end = period["start"][:7], period["end"][:7]
    year, month = int(start[:4]), int(start[5:7])
    while f"{year:04d}-{month:02d}" <= end:
        key = f"{year:04d}-{month:02d}"
        rows = [e for e in events if e["date"].startswith(key)]
        rows_posted = [e for e in rows if e["results_posted"]]
        rows_held = [e for e in rows_posted if e["trial_entries"] is not None]
        entries = sum(e["trial_entries"] for e in rows_held)
        tests = sum(e["test_entries"] or 0 for e in rows_posted)
        by_month.append({
            "month": key, "label": MONTH_LABELS[month - 1],
            "events": len(rows), "with_results": len(rows_posted),
            "trial_entries": entries, "test_entries": tests, "total_entries": entries + tests,
            "avg_trial_entries": _round(entries / len(rows_held)) if rows_held else None,
        })
        month += 1
        if month == 13:
            year, month = year + 1, 1

    trials = [{
        "event_number": e["event_number"], "date": e["date"], "club": e["club"], "city": e["city"],
        "state": e["state"], "scope": e["scope"], "trial_entries": e["trial_entries"],
        "trial_competitors": e["trial_competitors"], "test_entries": e["test_entries"],
    } for e in held]
    competition_rank(trials, "trial_entries")
    trials.sort(key=lambda t: (t["rank"], t["date"], t["event_number"]))

    def aggregate(key_field: str) -> list[dict]:
        groups: dict[str, dict] = {}
        for e in events:
            group = groups.setdefault(e[key_field], {
                key_field: e[key_field], "events": 0, "with_results": 0, "trials": 0,
                "trial_entries": 0, "largest_trial": None, "test_entries": 0, "states": set(),
            })
            group["events"] += 1
            group["states"].add(e["state"])
            if not e["results_posted"]:
                continue
            group["with_results"] += 1
            group["test_entries"] += e["test_entries"] or 0
            if e["trial_entries"] is not None:
                group["trials"] += 1
                group["trial_entries"] += e["trial_entries"]
                group["largest_trial"] = max(group["largest_trial"] or 0, e["trial_entries"])
        rows = []
        for group in groups.values():
            group["states"] = sorted(group["states"])
            group["avg_trial_entries"] = (
                _round(group["trial_entries"] / group["trials"]) if group["trials"] else None)
            # Only a club or state with a posted trial has a ranking figure.
            group["_rank_on"] = group["trial_entries"] if group["trials"] else None
            rows.append(group)
        competition_rank(rows, "_rank_on")
        for row in rows:
            del row["_rank_on"]
        rows.sort(key=lambda r: (r["rank"] is None, r["rank"] or 0, -r["events"], r[key_field]))
        return rows

    clubs = aggregate("club")
    for club in clubs:
        del club["events"], club["with_results"]
    states = aggregate("state")
    for state in states:
        del state["states"]

    feed = {
        "org": ORG,
        "name": "AKC",
        "program": "lure coursing",
        "season": season,
        "collected": collected,
        "period": period,
        "generated": today(),
        "source": {
            "name": SOURCE_NAME, "url": SOURCE_URL, "sheet_url": SHEET_URL,
            "file": source_file, "sha256": digest, "workbook_modified": parsed["modified"],
        },
        "stats": stats,
        "by_month": by_month,
        "rankings": {"trials": trials, "clubs": clubs, "states": states},
        "events": events,
        "notes": parsed["notes"],
    }
    assert_matches_sheet(feed, parsed)
    return feed


def assert_matches_sheet(feed: dict, parsed: dict) -> None:
    """Every figure the site shows must equal the sheet's own formulas."""
    problems: list[str] = []
    stats = feed["stats"]
    for key, sheet_value in parsed["summary"]["figures"].items():
        ours = stats.get(key)
        if key == "avg_trial_entries":
            same = ours is not None and sheet_value is not None and abs(ours - sheet_value) < 0.001
        else:
            same = ours == (int(sheet_value) if sheet_value is not None else None)
        if not same:
            problems.append(f"Summary {key}: sheet {sheet_value}, rows {ours}")
    sheet_months = parsed["summary"]["by_month"]
    for row in feed["by_month"]:
        sheet_row = sheet_months.get(row["month"])
        if sheet_row is None:
            problems.append(f"by-month {row['month']}: not on the Summary tab")
            continue
        for key in ("events", "with_results", "trial_entries", "test_entries", "total_entries"):
            if sheet_row[key] != row[key]:
                problems.append(f"by-month {row['month']} {key}: sheet {sheet_row[key]}, rows {row[key]}")
        sheet_avg, our_avg = sheet_row["avg_trial_entries"], row["avg_trial_entries"]
        if (sheet_avg is None) != (our_avg is None) or (sheet_avg is not None and abs(sheet_avg - our_avg) > 0.001):
            problems.append(f"by-month {row['month']} average: sheet {sheet_avg}, rows {our_avg}")
    covered = {row["month"] for row in feed["by_month"]}
    for key, sheet_row in sheet_months.items():
        if key not in covered and sheet_row["events"]:
            problems.append(f"by-month {key}: the sheet counts {sheet_row['events']} events outside the stated period")

    def compare(label: str, sheet_rows: list[dict], ours: dict[str, dict], key_of, fields) -> None:
        if len(sheet_rows) != len(ours):
            problems.append(f"{label}: sheet lists {len(sheet_rows)} rows, rows give {len(ours)}")
        for sheet_row in sheet_rows:
            key = key_of(sheet_row)
            mine = ours.get(key)
            if mine is None:
                problems.append(f"{label}: {key} is on the sheet but not in the rows")
                continue
            for field in fields:
                a, b = sheet_row[field], mine[field]
                if field == "avg_trial_entries":
                    same = (a is None and b is None) or (a is not None and b is not None and abs(a - b) < 0.001)
                else:
                    same = a == b
                if not same:
                    problems.append(f"{label} {key} {field}: sheet {a}, rows {b}")

    rankings = feed["rankings"]
    compare("Trial ranking", parsed["rankings"]["trials"],
            {(t["date"], t["club"], t["trial_entries"]): t for t in rankings["trials"]},
            lambda r: (r["date"], r["club"], r["trial_entries"]),
            ("rank", "city", "state", "trial_competitors", "test_entries"))
    compare("Club ranking", parsed["rankings"]["clubs"],
            {c["club"]: c for c in rankings["clubs"] if c["rank"] is not None},
            lambda r: r["club"],
            ("rank", "trials", "trial_entries", "avg_trial_entries", "largest_trial", "test_entries"))
    compare("State ranking", parsed["rankings"]["states"],
            {s["state"]: s for s in rankings["states"] if s["rank"] is not None},
            lambda r: r["state"],
            ("rank", "trials", "trial_entries", "avg_trial_entries", "largest_trial", "test_entries"))
    if problems:
        raise RacingParseError("the rows and the sheet's own figures disagree:\n  " + "\n  ".join(problems[:20]))


# ----------------------------------------------------------------- archive
def newest_raw() -> Path | None:
    files = sorted(RAW_DIR.glob("*.xlsx"))
    return files[-1] if files else None


def archive(raw: bytes, parsed: dict, collected: str, force: bool = False) -> Path:
    """File the export under its collected date. An unchanged re-export of
    the same collection is not written again (its zip bytes differ even
    when nothing in it changed); a changed one under the same date replaces
    the file with a notice, as the racing snapshots do."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    target = RAW_DIR / f"{collected}.xlsx"
    previous = newest_raw()
    if previous is not None and not force:
        kept = parse_workbook(previous.read_bytes())
        if content_key(kept) == content_key(parsed):
            print(f"Unchanged since {previous.name} - nothing archived.")
            return previous
    if target.exists():
        print(f"{target.name}: the collection was revised under the same date - replacing it.")
    target.write_bytes(raw)
    print(f"Archived {target.name}  {len(parsed['events'])} events")
    return target


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--file", type=Path, help="parse this export instead of fetching")
    parser.add_argument("--date", type=lambda text: date.fromisoformat(text).isoformat(),
                        help="the collected date, overriding the Summary tab")
    parser.add_argument("--offline", action="store_true",
                        help="rebuild from the newest archived export without fetching")
    parser.add_argument("--force", action="store_true", help="archive even if unchanged")
    args = parser.parse_args()

    raw: bytes | None = None
    if args.offline:
        source = newest_raw()
        if source is None:
            print("No archived export under data/akc/raw; fetch first.", file=sys.stderr)
            return 1
        raw = source.read_bytes()
    elif args.file:
        raw = args.file.read_bytes()
    else:
        try:
            raw = fetch()
        except (requests.RequestException, RacingParseError) as error:
            source = newest_raw()
            if source is None:
                print(f"Could not fetch the AKC workbook and nothing is archived: {error}", file=sys.stderr)
                return 1
            print(f"Could not fetch the AKC workbook; rebuilding from {source.name}. {error}",
                  file=sys.stderr)
            raw = source.read_bytes()
            args.offline = True

    parsed = parse_workbook(raw)
    collected = args.date or parsed["summary"]["collected"]
    if args.offline:
        source_path = newest_raw()
    else:
        source_path = archive(raw, parsed, collected, force=args.force)
    feed = build(parsed, source_path.name, sha256(source_path.read_bytes()), collected)
    size = write_json(FEED, feed)
    stats = feed["stats"]
    print(
        f"AKC entries collected {feed['collected']}: {stats['events']} events "
        f"({stats['with_results']} with results), {stats['trial_entries']:,} LC trial entries, "
        f"{stats['test_entries']:,} test entries, average trial {stats['avg_trial_entries']}, "
        f"{stats['clubs']} clubs in {stats['states']} states; feed {size:,} bytes"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
